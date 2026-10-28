"""Convert a Markdown file to .docx (headings, paragraphs, bullet/numbered lists, tables, code blocks).

Usage: python md_to_docx.py input.md output.docx
Requires python-docx (pip install python-docx).
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            run = paragraph.add_run(m.group(1) if m else part)
            run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        else:
            paragraph.add_run(part)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, r in enumerate(cells):
        row = table.add_row().cells
        for j in range(ncols):
            p = row[j].paragraphs[0]
            add_inline(p, r[j] if j < len(r) else "")
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def convert(src, dst):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            p.paragraph_format.left_indent = Pt(12)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                h = doc.add_heading(text, 0)
                h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                doc.add_heading(text, level - 1)
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, line[2:])
            i += 1
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Bullet 2" if m.group(1) else "List Bullet")
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^\s*\d+\.\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue
        if line.strip() in ("", "---"):
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print("wrote", sys.argv[2])
